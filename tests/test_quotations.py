import io
import unittest

from openpyxl import Workbook

from app.services.quotation_service import read_quotations


class QuotationTests(unittest.TestCase):
    def content(self, missing_price=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Bolo de teste'
        for index, (label, value) in enumerate([
            ('Custo Matéria-prima', 10), ('Custo Embalagem', 2),
            ('Custo fixo p/ Hora', 3), ('Custos variáveis', 1),
            ('Margem de Contribuição / Valor de venda', None if missing_price else 25),
            ('% Lucro liquido', 0.36), ('Rendimento', 1),
        ], 1):
            sheet.cell(index, 1, label)
            sheet.cell(index, 5, value)
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def test_costs_and_stable_link(self):
        first = read_quotations(self.content())
        row = first['rows'][0]
        self.assertEqual(row['cost_price'], '10.00')
        self.assertEqual(row['additional_cost'], '6.00')
        self.assertEqual(row['target_margin_percentage'], '36.00')
        self.assertEqual(row['sku'], read_quotations(self.content())['rows'][0]['sku'])

    def test_missing_calculated_price_is_not_imported(self):
        result = read_quotations(self.content(missing_price=True))
        self.assertEqual(result['rows'], [])
        self.assertEqual(len(result['warnings']), 1)

    def test_html_is_rejected(self):
        with self.assertRaises(ValueError):
            read_quotations(b'<html>login</html>')
