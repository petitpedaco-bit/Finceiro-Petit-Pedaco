"""Fichas editáveis com cálculo restrito a operações e funções de planilha.

Não executa código Python nem instruções contidas nos documentos.
"""
import ast
import io
import json
import math
import re
import posixpath
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.styles.numbers import BUILTIN_FORMATS
from pydantic import BaseModel, ConfigDict, Field, model_validator


class SheetCell(BaseModel):
    model_config = ConfigDict(allow_inf_nan=False)
    value: str | float | int | bool | None = None
    formula: str | None = Field(default=None, max_length=2000)
    format: str = Field(default='General', max_length=200)
    error: str | None = Field(default=None, max_length=200)


class QuotationSheet(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    rows: list[list[SheetCell]] = Field(max_length=2000)

    @model_validator(mode='after')
    def validate_columns(self):
        if any(len(row) > 100 for row in self.rows):
            raise ValueError('Cada aba pode ter até 100 colunas')
        return self


class QuotationBook(BaseModel):
    sheets: list[QuotationSheet] = Field(min_length=1, max_length=250)

    @model_validator(mode='after')
    def validate_size(self):
        if len({sheet.name for sheet in self.sheets}) != len(self.sheets):
            raise ValueError('Nomes de abas precisam ser únicos')
        if sum(len(row) for sheet in self.sheets for row in sheet.rows) > 200000:
            raise ValueError('Planilha excede 200 mil células')
        return self


def read_workbook(content: bytes) -> dict:
    """Lê células e fórmulas diretamente do Excel, incluindo fórmulas compartilhadas."""
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('Arquivo excede 20 MB')
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Planilha expandida excede 100 MB')
            strings = []
            if 'xl/sharedStrings.xml' in archive.namelist():
                strings = [''.join(item.itertext()) for item in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
            formats = dict(BUILTIN_FORMATS)
            style_ids = [0]
            if 'xl/styles.xml' in archive.namelist():
                styles = ET.fromstring(archive.read('xl/styles.xml'))
                for item in styles.findall(f'{ns}numFmts/{ns}numFmt'):
                    formats[int(item.get('numFmtId'))] = item.get('formatCode','General')
                style_ids = [int(item.get('numFmtId','0')) for item in styles.findall(f'{ns}cellXfs/{ns}xf')]
            relationships = {item.get('Id'): item.get('Target') for item in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
            root = ET.fromstring(archive.read('xl/workbook.xml'))
            sheets = []
            total_cells = 0
            for sheet in root.findall(f'{ns}sheets/{ns}sheet'):
                name = sheet.get('name')
                relation = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                target = relationships[relation]
                path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
                xml = ET.fromstring(archive.read(path))
                filled = {}; shared = {}; max_row = max_col = 1
                for cell in xml.findall(f'{ns}sheetData/{ns}row/{ns}c'):
                    coordinate = cell.get('r')
                    formula_node = cell.find(f'{ns}f')
                    raw = cell.findtext(f'{ns}v')
                    inline = cell.find(f'{ns}is')
                    if raw is None and inline is None and formula_node is None:
                        continue
                    column, row = coordinate_from_string(coordinate)
                    column = column_index_from_string(column)
                    if row > 2000 or column > 100:
                        raise ValueError(f'{name}: excede 2000 linhas ou 100 colunas')
                    formula = None
                    if formula_node is not None:
                        if formula_node.text:
                            formula = '=' + formula_node.text
                            if formula_node.get('t') == 'shared':
                                shared[formula_node.get('si')] = (formula,coordinate)
                        elif formula_node.get('t') == 'shared':
                            original, origin = shared[formula_node.get('si')]
                            formula = Translator(original, origin=origin).translate_formula(coordinate)
                    kind = cell.get('t')
                    value = None
                    if kind == 's' and raw is not None: value = strings[int(raw)]
                    elif kind == 'inlineStr' and inline is not None: value = ''.join(inline.itertext())
                    elif kind == 'b': value = raw == '1'
                    elif kind in ('str','e','d'): value = raw
                    elif raw not in (None, ''): value = float(raw)
                    style = int(cell.get('s','0'))
                    fmt = formats.get(style_ids[style] if style<len(style_ids) else 0,'General')
                    filled[(row,column)] = {'value':value,'formula':formula,'format':fmt}
                    max_row, max_col = max(max_row,row), max(max_col,column)
                total_cells += max_row * max_col
                if total_cells > 200000:
                    raise ValueError('Planilha excede 200 mil células')
                rows = [[filled.get((row,column), {'value':None,'formula':None,'format':'General'}) for column in range(1,max_col+1)] for row in range(1,max_row+1)]
                sheets.append({'name':name,'rows':rows})
            return QuotationBook(sheets=sheets).model_dump(mode='json')
    except (BadZipFile, KeyError, IndexError, ET.ParseError, TranslatorError) as exc:
        raise ValueError('Arquivo Excel inválido ou incompleto') from exc


class CellRange:
    def __init__(self, calculator, sheet_name, bounds):
        self.calculator, self.sheet_name, self.bounds = calculator, sheet_name, bounds

    def values(self):
        min_col, min_row, max_col, max_row = self.bounds
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                yield self.calculator.cell(self.sheet_name, row, column)

    def lookup(self, key, column):
        min_col, min_row, max_col, max_row = self.bounds
        if column < 1 or min_col + column - 1 > max_col:
            raise ValueError('Coluna de PROCV fora do intervalo')
        for row in range(min_row, max_row + 1):
            value = self.calculator.cell(self.sheet_name, row, min_col)
            if str(value).casefold() == str(key).casefold():
                return self.calculator.cell(self.sheet_name, row, min_col + column - 1)
        raise ValueError('Item não encontrado em fornecedores')


class WorkbookCalculator:
    def __init__(self, book: QuotationBook):
        self.book = book
        self.sheets = {sheet.name: sheet for sheet in book.sheets}
        self.memo = {}
        self.active = set()

    def cell(self, sheet_name, row, column):
        key = (sheet_name, row, column)
        if key in self.memo:
            value = self.memo[key]
            if isinstance(value, Exception):
                raise ValueError(str(value))
            return value
        if key in self.active or len(self.active) > 100:
            raise ValueError('Referência circular ou fórmula muito profunda')
        sheet = self.sheets.get(sheet_name)
        if sheet is None:
            raise ValueError('Aba referenciada não encontrada')
        if row < 1 or column < 1 or row > len(sheet.rows) or column > len(sheet.rows[row-1]):
            return None
        cell = sheet.rows[row-1][column-1]
        self.active.add(key)
        try:
            if not cell.formula and isinstance(cell.value,str) and cell.value in {
                '#REF!', '#DIV/0!', '#NAME?', '#N/A', '#VALUE!', '#NUM!', '#NULL!',
            }:
                raise ValueError(cell.value)
            value = self.expression(cell.formula, sheet_name) if cell.formula else cell.value
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError('Resultado numérico inválido')
            self.memo[key] = value
            return value
        except (ValueError, TypeError, ArithmeticError, SyntaxError, IndexError, RecursionError, TokenizerError) as exc:
            self.memo[key] = ValueError(str(exc)[:200])
            raise ValueError(str(exc)[:200]) from exc
        finally:
            self.active.remove(key)

    def reference(self, reference, sheet_name):
        if '!' in reference:
            sheet_name, reference = reference.rsplit('!', 1)
            sheet_name = sheet_name.strip("'").replace("''", "'")
        if '[' in sheet_name or sheet_name not in self.sheets:
            raise ValueError('Referência externa não suportada')
        if ':' not in reference and not re.fullmatch(r'\$?[A-Z]+\$?\d+', reference):
            raise ValueError('Referência ou nome definido não suportado')
        bounds = range_boundaries(reference.replace('$', ''))
        min_col, min_row, max_col, max_row = bounds
        sheet = self.sheets[sheet_name]
        min_col = min_col or 1
        max_col = max_col or max((len(row) for row in sheet.rows), default=1)
        min_row = min_row or 1
        max_row = min(max_row or len(sheet.rows), len(sheet.rows))
        if ':' not in reference:
            return self.cell(sheet_name, min_row, min_col)
        if max_col > 100 or max_row > 2000:
            raise ValueError('Intervalo excede limites')
        return CellRange(self, sheet_name, (min_col, min_row, max_col, max_row))

    @staticmethod
    def flatten(value):
        if isinstance(value, CellRange):
            return list(value.values())
        if isinstance(value, list):
            return [item for part in value for item in WorkbookCalculator.flatten(part)]
        return [value]

    def expression(self, formula, sheet_name):
        pieces, references = [], {}
        for token in Tokenizer(formula).items:
            if token.type == 'OPERAND' and token.subtype == 'RANGE':
                name = f'ref{len(references)}'
                references[name] = self.reference(token.value, sheet_name)
                pieces.append(name)
            elif token.type == 'OPERAND' and token.subtype == 'TEXT':
                pieces.append(json.dumps(token.value[1:-1].replace('""', '"')))
            elif token.type == 'OPERAND' and token.subtype == 'LOGICAL':
                pieces.append('True' if token.value.upper() == 'TRUE' else 'False')
            elif token.type == 'OPERAND' and token.subtype == 'ERROR':
                raise ValueError(token.value)
            elif token.value == '%':
                pieces.append('/100')
            else:
                pieces.append(token.value.replace('^', '**').replace(';', ','))
        tree = ast.parse(''.join(pieces), mode='eval')
        def evaluate(node):
            if isinstance(node, ast.Constant):
                return node.value
            if isinstance(node, ast.Name) and node.id in references:
                return references[node.id]
            if isinstance(node, ast.UnaryOp):
                value = evaluate(node.operand) or 0
                if isinstance(node.op, ast.USub): return -value
                if isinstance(node.op, ast.UAdd): return value
            if isinstance(node, ast.BinOp):
                left, right = evaluate(node.left) or 0, evaluate(node.right) or 0
                if isinstance(node.op, ast.Add): return left + right
                if isinstance(node.op, ast.Sub): return left - right
                if isinstance(node.op, ast.Mult): return left * right
                if isinstance(node.op, ast.Div): return left / right
                if isinstance(node.op, ast.Pow) and abs(right) <= 10: return left ** right
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
                args = [evaluate(arg) for arg in node.args]
                name = node.func.id.upper()
                if name in {'SUM', 'MIN', 'MAX'}:
                    numbers = [v for arg in args for v in self.flatten(arg) if isinstance(v, (int, float))]
                    if name == 'SUM': return sum(numbers)
                    if name == 'MIN': return min(numbers)
                    return max(numbers)
                if name == 'ROUND':
                    from decimal import Decimal, ROUND_HALF_UP
                    digits = int(args[1]) if len(args)>1 else 0
                    if abs(digits) > 20: raise ValueError('Precisão de arredondamento excede limites')
                    return float(Decimal(str(args[0])).quantize(Decimal(1).scaleb(-digits),rounding=ROUND_HALF_UP))
                if name == 'VLOOKUP':
                    if len(args) < 4 or args[3] not in (False, 0):
                        raise ValueError('PROCV aproximado precisa de revisão')
                    if not isinstance(args[1], CellRange):
                        raise ValueError('PROCV exige um intervalo')
                    return args[1].lookup(args[0], int(args[2]))
            raise ValueError('Fórmula não suportada: revise antes de vincular ao produto')
        return evaluate(tree.body)

    def calculate(self):
        warnings = []
        for sheet in self.book.sheets:
            for row_index, row in enumerate(sheet.rows, 1):
                for column_index, cell in enumerate(row, 1):
                    cell.error = None
                    if cell.formula:
                        try:
                            cell.value = self.cell(sheet.name, row_index, column_index)
                        except ValueError as exc:
                            cell.value = None
                            cell.error = str(exc)[:200]
                            warnings.append(f'{sheet.name}, linha {row_index}: {cell.error}')
        return warnings


def workbook_quotations(book: QuotationBook) -> dict:
    """Extrai os resultados recalculados sem reconstruir um arquivo Excel."""
    import hashlib
    from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
    from app.services.quotation_service import QuotationRow, normalize
    rows, warnings = [], []
    for sheet in book.sheets:
        fields = {}
        for row in sheet.rows:
            if len(row) < 5: continue
            label = normalize(row[0].value)
            for key, marker in (
                ('materials','custo materia-prima'),('packaging','custo embalagem'),
                ('fixed','custo fixo p/'),('variable','custos variaveis'),
                ('price','margem de contribuicao / valor de venda'),
                ('margin','% lucro liquido'),('yield','rendimento'),
            ):
                if marker in label: fields[key] = row[4].value
        if 'materials' not in fields: continue
        try:
            def number(key):
                value = fields.get(key)
                if value is None or isinstance(value,bool): raise ValueError(f'{key}: valor calculado ausente')
                value = Decimal(str(value))
                if not value.is_finite() or value < 0: raise ValueError(f'{key}: valor inválido')
                return value
            def money(value): return value.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
            if number('yield') <= 0: raise ValueError('Rendimento deve ser positivo')
            rows.append(QuotationRow(sheet=sheet.name,name=sheet.name,
                sku='COT-'+hashlib.sha256(sheet.name.encode()).hexdigest()[:16].upper(),
                cost_price=money(number('materials')),
                additional_cost=money(sum((number(key) for key in ('fixed','packaging','variable')),Decimal(0))),
                sale_price=money(number('price')),target_margin_percentage=money(number('margin')*100)).model_dump(mode='json'))
        except (ValueError,InvalidOperation) as exc:
            warnings.append(f'{sheet.name}: {exc}')
    return {'rows':rows,'warnings':warnings}
