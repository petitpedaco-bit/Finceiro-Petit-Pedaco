"""Lê resultados calculados das fichas de cotação sem executar fórmulas."""
import hashlib
import io
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from pydantic import BaseModel, Field


class QuotationRow(BaseModel):
    sheet: str
    sku: str = Field(min_length=1, max_length=64)
    name: str = Field(min_length=1, max_length=160)
    cost_price: Decimal = Field(ge=0, max_digits=14, decimal_places=2)
    additional_cost: Decimal = Field(ge=0, max_digits=14, decimal_places=2)
    sale_price: Decimal = Field(gt=0, max_digits=14, decimal_places=2)
    target_margin_percentage: Decimal = Field(ge=0, lt=100)


def normalize(value: object) -> str:
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').lower()) if not unicodedata.combining(c))


def read_quotations(content: bytes) -> dict:
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('Arquivo excede 20 MB')
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Planilha expandida excede 100 MB')
    except BadZipFile as exc:
        raise ValueError('Resposta não é um Excel válido. Confira o compartilhamento público.') from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows, warnings = [], []
    try:
        for sheet in workbook.worksheets:
            fields = {}
            for cells in sheet.iter_rows(max_row=min(sheet.max_row or 500, 500), max_col=5, values_only=True):
                label = normalize(cells[0])
                for key, marker in (
                    ('materials', 'custo materia-prima'), ('packaging', 'custo embalagem'),
                    ('fixed', 'custo fixo p/'), ('variable', 'custos variaveis'),
                    ('price', 'margem de contribuicao / valor de venda'),
                    ('margin', '% lucro liquido'), ('yield', 'rendimento'),
                ):
                    if marker in label:
                        fields[key] = cells[4]
            if 'materials' not in fields:
                continue
            try:
                def number(key: str) -> Decimal:
                    value = fields.get(key)
                    if value is None or isinstance(value, bool):
                        raise ValueError(f'{key}: valor calculado ausente')
                    parsed = Decimal(str(value))
                    if not parsed.is_finite() or parsed < 0:
                        raise ValueError(f'{key}: valor inválido')
                    return parsed
                unit = Decimal('0.01')
                cost = number('materials').quantize(unit, rounding=ROUND_HALF_UP)
                extras = sum((number(key) for key in ('fixed', 'packaging', 'variable')), Decimal(0)).quantize(unit, rounding=ROUND_HALF_UP)
                price = number('price').quantize(unit, rounding=ROUND_HALF_UP)
                margin = (number('margin') * 100).quantize(unit)
                if number('yield') <= 0:
                    raise ValueError('rendimento precisa ser positivo')
                # SKU determinístico por aba; o usuário pode informar o SKU já cadastrado.
                sku = 'COT-' + hashlib.sha256(sheet.title.encode()).hexdigest()[:16].upper()
                rows.append(QuotationRow(sheet=sheet.title, sku=sku, name=sheet.title.strip(),
                    cost_price=cost, additional_cost=extras, sale_price=price,
                    target_margin_percentage=margin).model_dump(mode='json'))
            except (ValueError, InvalidOperation) as exc:
                warnings.append(f'{sheet.title}: {exc}')
    finally:
        workbook.close()
    return {'rows': rows, 'warnings': warnings,
            'notice': 'Valores calculados salvos no Excel. Recalcule e salve a planilha antes de importar. Taxa de cartão não integra o custo importado.'}
